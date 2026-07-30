"""Load local bibliographic records from CSV, TSV, or Excel (.xlsx) exports.

Built around the item-level export shape (Title, Author, Format, Pub Date,
Call No., Sort Title, Sort Author) but accepts any column layout via a
mapping. Identifier columns (ISBN, ISSN, OCLC, LCCN), when present, are
carried along but treated as *hints only* — per project background, they are
unreliable in this catalog.
"""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, Iterator

_YEAR = re.compile(r"(?<!\d)(1[4-9]\d{2}|20\d{2})(?!\d)")  # tolerates "c1976", "[1897?]"
_SERIAL_WORDS = re.compile(r"serial|periodical|journal|magazine|newspaper|continuing", re.IGNORECASE)

# Canonical field names; --csv-map overrides ("id=BibID,title=Title245").
# Only title is required — records without an id column get stable row-number
# ids ("row-2" = spreadsheet row 2). Multi-value cells split on ";".
CSV_FIELDS = (
    "id", "title", "author", "year", "publisher", "material_type",
    "call_number", "isbn", "issn", "oclc", "lccn",
)

# Header spellings recognized without any --csv-map, matched case-insensitively.
# Covers both the canonical names and the ItemRecordExport.xlsx layout.
_DEFAULT_ALIASES = {
    "id": ("id", "record id", "bib id"),
    "title": ("title",),
    "author": ("author",),
    "year": ("year", "pub date", "publication date", "date"),
    "publisher": ("publisher",),
    "material_type": ("material_type", "format", "material type"),
    "call_number": ("call_number", "call no.", "call no", "call number"),
    "isbn": ("isbn",),
    "issn": ("issn",),
    "oclc": ("oclc", "oclc number", "oclc no."),
    "lccn": ("lccn",),
}


@dataclass
class LocalRecord:
    record_id: str
    title: str
    author: str = ""
    year: str = ""
    publisher: str = ""
    material_type: str = ""
    call_number: str = ""
    # Untrusted identifier hints from the source record
    isbns: list[str] = field(default_factory=list)
    issns: list[str] = field(default_factory=list)
    oclc_numbers: list[str] = field(default_factory=list)
    lccn: str = ""

    @property
    def is_serial(self) -> bool:
        return bool(_SERIAL_WORDS.search(self.material_type))


def load_records(path: str | Path, csv_map: dict[str, str] | None = None) -> Iterator[LocalRecord]:
    """Yield LocalRecords from a .csv, .tsv, or .xlsx file."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in (".csv", ".tsv"):
        yield from _load_csv(path, csv_map or {}, delimiter="\t" if suffix == ".tsv" else ",")
    elif suffix == ".xlsx":
        yield from _load_xlsx(path, csv_map or {})
    else:
        raise ValueError(f"Unsupported input format: {path.name} (expected .csv, .tsv, or .xlsx)")


def _load_csv(path: Path, csv_map: dict[str, str], delimiter: str) -> Iterator[LocalRecord]:
    with open(path, newline="", encoding="utf-8-sig") as fh:
        reader = csv.DictReader(fh, delimiter=delimiter)
        if reader.fieldnames is None:
            return
        yield from _rows_to_records(reader, list(reader.fieldnames), csv_map)


def _load_xlsx(path: Path, csv_map: dict[str, str]) -> Iterator[LocalRecord]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            return
        columns = [str(h).strip() if h is not None else "" for h in header]
        dict_rows = (
            {col: ("" if cell is None else str(cell)) for col, cell in zip(columns, row)}
            for row in rows
        )
        yield from _rows_to_records(dict_rows, columns, csv_map)
    finally:
        workbook.close()


def _resolve_columns(columns: list[str], csv_map: dict[str, str]) -> dict[str, str]:
    """Map canonical field -> actual header, honoring csv_map then default aliases."""
    unknown = set(csv_map) - set(CSV_FIELDS)
    if unknown:
        raise ValueError(f"--csv-map has unknown field(s) {sorted(unknown)}. Known: {', '.join(CSV_FIELDS)}")
    by_lower = {col.lower(): col for col in columns}
    colmap: dict[str, str] = {}
    for canon in CSV_FIELDS:
        if canon in csv_map:
            if csv_map[canon] not in columns:
                raise ValueError(f"--csv-map points '{canon}' at column '{csv_map[canon]}', which is not in the input. Available: {columns}")
            colmap[canon] = csv_map[canon]
            continue
        for alias in _DEFAULT_ALIASES[canon]:
            if alias in by_lower:
                colmap[canon] = by_lower[alias]
                break
    if "title" not in colmap:
        raise ValueError(f"Could not find a title column in {columns}. Use --csv-map title=YourColumn.")
    return colmap


def _rows_to_records(
    rows: Iterable[dict], columns: list[str], csv_map: dict[str, str]
) -> Iterator[LocalRecord]:
    colmap = _resolve_columns(columns, csv_map)

    def get(row: dict, canon: str) -> str:
        if canon not in colmap:
            return ""
        return (row.get(colmap[canon]) or "").strip()

    def multi(row: dict, canon: str) -> list[str]:
        return [part.strip() for part in get(row, canon).split(";") if part.strip()]

    for row_number, row in enumerate(rows, start=2):  # row 1 is the header
        if not get(row, "title"):
            continue  # blank spreadsheet row
        yield LocalRecord(
            record_id=get(row, "id") or f"row-{row_number}",
            title=get(row, "title"),
            author=get(row, "author"),
            year=(m.group(1) if (m := _YEAR.search(get(row, "year"))) else ""),
            publisher=get(row, "publisher"),
            material_type=get(row, "material_type"),
            call_number=get(row, "call_number"),
            isbns=[re.sub(r"[^0-9Xx]", "", part).upper() for part in multi(row, "isbn") if re.sub(r"[^0-9Xx]", "", part)],
            issns=[part.upper().replace(" ", "") for part in multi(row, "issn")],
            oclc_numbers=[digits.lstrip("0") or "0" for part in multi(row, "oclc") if (digits := re.sub(r"\D", "", part))],
            lccn=get(row, "lccn"),
        )
